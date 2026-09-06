"""
advanced.py
-----------
The "15 must-have features" module. Kept separate from main.py so the core
upload/clean/analyze/chat pipeline stays easy to read on its own.

Every endpoint here is defensive about missing state (no dataset uploaded,
analysis not run yet, GROQ_API_KEY not set, etc.) — features degrade
gracefully to heuristic output instead of crashing when the optional LLM
isn't available, since these are meant to work in a plain demo/hackathon
deploy with no paid keys required (except AI chat, which already needs one).
"""

import io
import re
from datetime import datetime

import numpy as np
import pandas as pd
from fastapi import APIRouter, HTTPException
from fastapi.responses import Response
from pydantic import BaseModel

from state import STATE

router = APIRouter()


def _require_clean_df() -> pd.DataFrame:
    if STATE["clean_df"] is None:
        raise HTTPException(status_code=400, detail="No dataset uploaded/cleaned yet.")
    return STATE["clean_df"]


def _require_ml_result() -> dict:
    if STATE["last_ml_result"] is None:
        raise HTTPException(status_code=400, detail="Run 'Start AI Analysis' first.")
    return STATE["last_ml_result"]


# ---------------------------------------------------------------------------
# 6. AI Dataset Health Score
# ---------------------------------------------------------------------------
def compute_health_score(df: pd.DataFrame) -> dict:
    total_cells = df.shape[0] * df.shape[1] or 1
    missing_pct = round(float(df.isna().sum().sum()) / total_cells * 100, 2)
    dup_pct = round(float(df.duplicated().sum()) / max(len(df), 1) * 100, 2)
    null_columns = df.columns[df.isna().all()].tolist()

    outlier_hits, outlier_checked = 0, 0
    for col in df.select_dtypes(include=[np.number]).columns[:15]:
        if df[col].nunique() <= 10:
            continue
        q1, q3 = df[col].quantile(0.25), df[col].quantile(0.75)
        iqr = q3 - q1
        lower, upper = q1 - 1.5 * iqr, q3 + 1.5 * iqr
        outlier_hits += int(((df[col] < lower) | (df[col] > upper)).sum())
        outlier_checked += len(df)
    outlier_pct = round(outlier_hits / outlier_checked * 100, 2) if outlier_checked else 0.0

    score = 100.0
    score -= min(30, missing_pct * 1.5)
    score -= min(25, dup_pct * 2)
    score -= min(20, outlier_pct * 1.2)
    score -= min(15, len(null_columns) * 5)
    score = round(max(0.0, min(100.0, score)), 1)

    if score >= 90:
        grade = "Excellent"
    elif score >= 75:
        grade = "Good"
    elif score >= 55:
        grade = "Fair"
    else:
        grade = "Needs Attention"

    return {
        "score": score,
        "grade": grade,
        "breakdown": {
            "missing_values_pct": missing_pct,
            "duplicate_rows_pct": dup_pct,
            "outlier_rate_pct": outlier_pct,
            "empty_columns": null_columns,
        },
    }


@router.get("/api/health-score")
async def health_score():
    if STATE["raw_df"] is None:
        raise HTTPException(status_code=400, detail="No dataset uploaded yet.")
    return compute_health_score(STATE["raw_df"])


# ---------------------------------------------------------------------------
# Shared builder used by both /api/ai-consultant and /api/presentation so the
# two features never disagree with each other.
# ---------------------------------------------------------------------------
PERSONA_NAME_POOL = [
    ("Aarav", "👨"), ("Priya", "👩"), ("Rohan", "🧑"), ("Sneha", "👩"),
    ("Karan", "👨"), ("Isha", "👩"), ("Vikram", "🧑"), ("Meera", "👩"),
]


def _find_column_like(columns, keywords):
    for col in columns:
        low = str(col).lower()
        if any(k in low for k in keywords):
            return col
    return None


def _build_consultant_data() -> dict:
    df = _require_clean_df()
    ml = _require_ml_result()
    health = compute_health_score(STATE["raw_df"] if STATE["raw_df"] is not None else df)
    profiles = STATE["cluster_profiles"] or []

    insights = []
    top_profile = profiles[0] if profiles else None
    risk_profile = profiles[-1] if len(profiles) > 1 else None

    if top_profile:
        insights.append(
            f"The '{top_profile['tier']}' segment makes up {top_profile['pct_of_total']}% of records "
            f"and shows the strongest overall metrics — prioritize retention and upsell here."
        )
    if risk_profile:
        insights.append(
            f"{risk_profile['pct_of_total']}% of records fall into the '{risk_profile['tier']}' segment "
            f"({risk_profile['size']} rows) — a re-engagement campaign could reduce this risk."
        )
    if ml.get("feature_importance"):
        top_feat = ml["feature_importance"][0]
        insights.append(
            f"'{top_feat['feature']}' is the strongest predictor of {ml.get('target_column') or 'the target'}, "
            f"driving {round(top_feat['importance'] * 100, 1)}% of the model's decisions."
        )
    corr = ml.get("correlation")
    if corr and corr.get("columns") and len(corr["columns"]) >= 2:
        cols, matrix = corr["columns"], corr["matrix"]
        best_pair, best_val = None, 0
        for i in range(len(cols)):
            for j in range(i + 1, len(cols)):
                v = matrix[i][j]
                if abs(v) > abs(best_val):
                    best_val, best_pair = v, (cols[i], cols[j])
        if best_pair and abs(best_val) >= 0.4:
            direction = "positively" if best_val > 0 else "negatively"
            insights.append(
                f"'{best_pair[0]}' and '{best_pair[1]}' are {direction} correlated (r={best_val:.2f}) — "
                f"worth exploring for causality or combined targeting."
            )
    if health["score"] < 75:
        insights.append(
            f"Data quality score is {health['score']}% ({health['grade']}) — "
            f"{health['breakdown']['duplicate_rows_pct']}% duplicate rows and "
            f"{health['breakdown']['missing_values_pct']}% missing cells were found in the raw upload."
        )
    if ml.get("best_model") and ml.get("best_model") != "None":
        acc = ml.get("best_accuracy", 0)
        quality_note = "strong enough to act on" if acc >= 80 else "a reasonable baseline — more data would help"
        insights.append(f"The {ml['best_model']} reached {acc}% accuracy — {quality_note}.")

    if not insights:
        insights.append("Upload a larger or richer dataset for deeper AI-generated insights.")
    insights = insights[:5]

    # --- Recommendation cards ---
    cards = []
    if top_profile:
        cards.append({
            "icon": "💰", "tag": "Revenue Opportunity",
            "title": f"Grow the {top_profile['tier']} segment",
            "body": f"{top_profile['pct_of_total']}% of customers already sit in your best-performing "
                    f"segment. Increasing marketing spend here typically compounds fastest.",
        })
    if risk_profile:
        cards.append({
            "icon": "⚠️", "tag": "High Risk",
            "title": f"{risk_profile['size']} customers need attention",
            "body": f"These fall in the '{risk_profile['tier']}' segment ({risk_profile['pct_of_total']}% "
                    f"of records). Consider a discount, check-in, or win-back campaign.",
        })
    if ml.get("feature_importance"):
        top_feat = ml["feature_importance"][0]
        cards.append({
            "icon": "🎯", "tag": "Key Driver",
            "title": f"'{top_feat['feature']}' drives outcomes",
            "body": f"This single column explains {round(top_feat['importance'] * 100, 1)}% of what the "
                    f"model uses to predict {ml.get('target_column') or 'the outcome'}. Track it closely.",
        })
    cards.append({
        "icon": "📊", "tag": "Data Quality",
        "title": f"{health['score']}% health score ({health['grade']})",
        "body": f"{health['breakdown']['missing_values_pct']}% missing cells, "
                f"{health['breakdown']['duplicate_rows_pct']}% duplicate rows, "
                f"{health['breakdown']['outlier_rate_pct']}% outlier rate before cleaning.",
    })

    # --- Executive summary (McKinsey-style scorecard) ---
    silhouette = ml.get("kmeans", {}).get("silhouette", 0) or 0
    silhouette_pct = max(0, min(100, (silhouette + 1) * 50))
    accuracy_component = ml.get("best_accuracy") or 60
    business_score = round(0.45 * health["score"] + 0.35 * accuracy_component + 0.20 * silhouette_pct)
    business_score = max(0, min(100, business_score))

    at_risk_pct = risk_profile["pct_of_total"] if risk_profile else 0
    revenue_risk = "High" if at_risk_pct > 30 else "Medium" if at_risk_pct > 15 else "Low"

    satisfaction_map = {"Excellent": "High", "Good": "High", "Fair": "Medium", "Needs Attention": "Low"}
    customer_satisfaction = satisfaction_map.get(health["grade"], "Medium")

    top_pct = top_profile["pct_of_total"] if top_profile else 0
    growth_opportunity = "Excellent" if top_pct > 40 else "Good" if top_pct > 20 else "Moderate"

    if risk_profile and top_profile:
        recommended_action = (
            f"Launch a loyalty campaign for the '{top_profile['tier']}' segment while running a "
            f"win-back offer for the '{risk_profile['tier']}' segment."
        )
    elif top_profile:
        recommended_action = f"Double down on the '{top_profile['tier']}' segment — it's your strongest performer."
    else:
        recommended_action = "Run AI Analysis with a larger dataset to unlock segment-specific recommendations."

    executive_summary = {
        "business_score": business_score,
        "revenue_risk": revenue_risk,
        "customer_satisfaction": customer_satisfaction,
        "growth_opportunity": growth_opportunity,
        "recommended_action": recommended_action,
    }

    # --- Personas ---
    personas = []
    age_col = _find_column_like(df.columns, ["age"])
    income_col = _find_column_like(df.columns, ["income", "salary", "revenue"])
    for i, p in enumerate(profiles[:4]):
        name, avatar = PERSONA_NAME_POOL[i % len(PERSONA_NAME_POOL)]
        row = p.get("representative_row", {}) or {}
        means = p.get("feature_means", {})
        # pick the 2 features where this cluster deviates most from the overall mean
        # (overall mean approximated by averaging all cluster means for that feature)
        traits = []
        if means:
            all_vals = {k: [] for k in means}
            for other in profiles:
                for k, v in other.get("feature_means", {}).items():
                    all_vals.setdefault(k, []).append(v)
            deviations = []
            for k, v in means.items():
                overall = np.mean(all_vals.get(k, [v]))
                if overall:
                    deviations.append((k, (v - overall) / (abs(overall) + 1e-9)))
            deviations.sort(key=lambda x: abs(x[1]), reverse=True)
            for k, dev in deviations[:2]:
                traits.append(f"{'High' if dev > 0 else 'Low'} {k}")
        personas.append({
            "name": name,
            "avatar": avatar,
            "age": row.get(age_col) if age_col else None,
            "income": row.get(income_col) if income_col else None,
            "traits": traits,
            "prediction": f"{p['tier']} Customer",
            "segment_pct": p["pct_of_total"],
        })

    return {
        "insights": insights,
        "recommendation_cards": cards[:4],
        "executive_summary": executive_summary,
        "health": health,
        "personas": personas,
        "generated_at": datetime.utcnow().isoformat(),
    }


@router.get("/api/ai-consultant")
async def ai_consultant():
    """Feature #1 (Business Consultant), #5 (Recommendation Cards),
    #7 (Persona Generator), #15 (Executive Summary) — all in one payload
    so the frontend can render them the instant analysis finishes, without
    the user asking a single question."""
    return _build_consultant_data()


# ---------------------------------------------------------------------------
# 8. AI What-If Simulator
# ---------------------------------------------------------------------------
@router.get("/api/whatif/schema")
async def whatif_schema():
    df = _require_clean_df()
    if not STATE["rf_feature_cols"] and not STATE["kmeans_feature_cols"]:
        raise HTTPException(status_code=400, detail="Run 'Start AI Analysis' first.")

    all_cols = list(dict.fromkeys(STATE["rf_feature_cols"] + STATE["kmeans_feature_cols"]))
    fields = []
    for col in all_cols:
        if col not in df.columns:
            continue
        if col in STATE["label_encoders"]:
            enc = STATE["label_encoders"][col]
            fields.append({"name": col, "type": "categorical", "options": [str(c) for c in enc.classes_]})
        elif pd.api.types.is_numeric_dtype(df[col]):
            fields.append({
                "name": col, "type": "numeric",
                "min": round(float(df[col].min()), 2),
                "max": round(float(df[col].max()), 2),
                "mean": round(float(df[col].mean()), 2),
            })
    return {
        "fields": fields,
        "target_column": STATE["target_col"],
        "has_classifier": STATE["trained_rf"] is not None,
        "has_clustering": STATE["trained_kmeans"] is not None,
    }


class WhatIfRequest(BaseModel):
    values: dict


@router.post("/api/whatif")
async def whatif(req: WhatIfRequest):
    if STATE["trained_rf"] is None and STATE["trained_kmeans"] is None:
        raise HTTPException(status_code=400, detail="Run 'Start AI Analysis' first.")

    result = {}

    if STATE["trained_rf"] is not None:
        row = {}
        for col in STATE["rf_feature_cols"]:
            raw_val = req.values.get(col, STATE["feature_defaults"].get(col, 0))
            if col in STATE["label_encoders"]:
                enc = STATE["label_encoders"][col]
                str_val = str(raw_val)
                if str_val in enc.classes_:
                    row[col] = enc.transform([str_val])[0]
                else:
                    row[col] = STATE["feature_defaults"].get(col, 0)
            else:
                try:
                    row[col] = float(raw_val)
                except (TypeError, ValueError):
                    row[col] = STATE["feature_defaults"].get(col, 0)
        X_row = pd.DataFrame([row])[STATE["rf_feature_cols"]]
        rf = STATE["trained_rf"]
        proba = rf.predict_proba(X_row)[0]
        classes = rf.classes_
        best_idx = int(np.argmax(proba))
        confidence = float(proba[best_idx])
        result["predicted_class"] = str(classes[best_idx])
        result["class_probability"] = round(confidence * 100, 1)
        result["target_column"] = STATE["target_col"]
        result["confidence_label"] = "High" if confidence >= 0.7 else "Medium" if confidence >= 0.45 else "Low"

    if STATE["trained_kmeans"] is not None:
        row = {}
        for col in STATE["kmeans_feature_cols"]:
            raw_val = req.values.get(col, STATE["feature_defaults"].get(col, 0))
            try:
                row[col] = float(raw_val)
            except (TypeError, ValueError):
                row[col] = STATE["feature_defaults"].get(col, 0)
        X_row = pd.DataFrame([row])[STATE["kmeans_feature_cols"]]
        cluster_id = int(STATE["trained_kmeans"].predict(X_row)[0])
        profile = next((p for p in STATE["cluster_profiles"] if p["cluster_id"] == cluster_id), None)
        result["cluster_id"] = cluster_id
        result["cluster_tier"] = profile["tier"] if profile else f"Cluster {cluster_id}"
        result["cluster_pct_of_total"] = profile["pct_of_total"] if profile else None

    return result


# ---------------------------------------------------------------------------
# 11. AI Smart Search
# ---------------------------------------------------------------------------
class SmartSearchRequest(BaseModel):
    query: str


@router.post("/api/smart-search")
async def smart_search(req: SmartSearchRequest):
    df = _require_clean_df()
    q = req.query.lower().strip()
    if not q:
        raise HTTPException(status_code=400, detail="Empty query.")

    col_lookup = {str(c).lower(): c for c in df.columns}
    target_col = None
    for low, orig in sorted(col_lookup.items(), key=lambda kv: -len(kv[0])):
        if low in q or low.replace("_", " ") in q:
            target_col = orig
            break

    result_df = df
    note = ""

    if target_col is not None and pd.api.types.is_numeric_dtype(df[target_col]):
        if any(w in q for w in ["high", "top", "above", "over", "greater", "more than"]):
            thresh = float(df[target_col].quantile(0.75))
            result_df = df[df[target_col] > thresh]
            note = f"Rows where '{target_col}' is above the 75th percentile ({thresh:.2f})"
        elif any(w in q for w in ["low", "bottom", "below", "under", "less than", "cheap"]):
            thresh = float(df[target_col].quantile(0.25))
            result_df = df[df[target_col] < thresh]
            note = f"Rows where '{target_col}' is below the 25th percentile ({thresh:.2f})"
        else:
            note = f"Matched numeric column '{target_col}' — showing all rows sorted descending"
            result_df = df.sort_values(target_col, ascending=False)
    elif target_col is not None:
        matched_val = None
        for val in df[target_col].astype(str).unique():
            if val.lower() and val.lower() in q:
                matched_val = val
                break
        if matched_val is not None:
            result_df = df[df[target_col].astype(str) == matched_val]
            note = f"Rows where '{target_col}' = '{matched_val}'"
        else:
            note = f"Matched column '{target_col}' but no specific value recognized — showing first rows"
    else:
        note = "Couldn't match a column name in that query — showing the first rows of the dataset"

    limit_match = re.search(r"top\s+(\d+)|first\s+(\d+)", q)
    limit = int(limit_match.group(1) or limit_match.group(2)) if limit_match else 25
    result_df = result_df.head(min(limit, 200))

    return {
        "note": note,
        "matched_rows": int(result_df.shape[0]),
        "columns": list(result_df.columns),
        "rows": result_df.replace({np.nan: None}).to_dict(orient="records"),
    }


# ---------------------------------------------------------------------------
# 10. AI Auto Presentation Generator (PowerPoint)
# ---------------------------------------------------------------------------
@router.get("/api/presentation")
async def generate_presentation():
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor
    except ImportError:
        raise HTTPException(
            status_code=500,
            detail="python-pptx isn't installed on the server. Add 'python-pptx' to requirements.txt.",
        )

    data = _build_consultant_data()
    ml = STATE["last_ml_result"] or {}
    filename = STATE["filename"] or "dataset"

    NAVY = RGBColor(0x14, 0x1E, 0x3C)
    TEAL = RGBColor(0x14, 0x64, 0x8C)
    GREY = RGBColor(0x5A, 0x5A, 0x5A)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    def add_slide():
        return prs.slides.add_slide(blank)

    def add_title(slide, text, size=32):
        box = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(12), Inches(1))
        tf = box.text_frame
        tf.text = text
        run = tf.paragraphs[0].runs[0]
        run.font.size = Pt(size)
        run.font.bold = True
        run.font.color.rgb = NAVY
        return box

    def add_bullets(slide, items, top=1.5, size=18):
        box = slide.shapes.add_textbox(Inches(0.8), Inches(top), Inches(11.5), Inches(5.5))
        tf = box.text_frame
        tf.word_wrap = True
        for i, item in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = f"•  {item}"
            p.font.size = Pt(size)
            p.font.color.rgb = GREY
            p.space_after = Pt(12)

    # Slide 1: Title
    s = add_slide()
    box = s.shapes.add_textbox(Inches(0.8), Inches(2.6), Inches(11.5), Inches(1.5))
    tf = box.text_frame
    tf.text = "AI Executive Summary"
    tf.paragraphs[0].runs[0].font.size = Pt(44)
    tf.paragraphs[0].runs[0].font.bold = True
    tf.paragraphs[0].runs[0].font.color.rgb = NAVY
    sub = s.shapes.add_textbox(Inches(0.8), Inches(3.6), Inches(11.5), Inches(1))
    sub.text_frame.text = f"Dataset: {filename}  ·  Generated {datetime.utcnow().strftime('%Y-%m-%d')}"
    sub.text_frame.paragraphs[0].runs[0].font.size = Pt(18)
    sub.text_frame.paragraphs[0].runs[0].font.color.rgb = TEAL

    # Slide 2: Business scorecard
    s = add_slide()
    add_title(s, "Business Scorecard")
    ex = data["executive_summary"]
    add_bullets(s, [
        f"Business Score: {ex['business_score']}/100",
        f"Revenue Risk: {ex['revenue_risk']}",
        f"Customer Satisfaction: {ex['customer_satisfaction']}",
        f"Growth Opportunity: {ex['growth_opportunity']}",
        f"Recommended Action: {ex['recommended_action']}",
    ])

    # Slide 3: Key insights
    s = add_slide()
    add_title(s, "Key AI Insights")
    add_bullets(s, data["insights"])

    # Slide 4: ML results
    s = add_slide()
    add_title(s, "Machine Learning Results")
    add_bullets(s, [
        f"K-Means Clusters Found: {ml.get('kmeans', {}).get('clusters', 0)}",
        f"Decision Tree Accuracy: {ml.get('decision_tree', {}).get('accuracy', 0)}%",
        f"Random Forest Accuracy: {ml.get('random_forest', {}).get('accuracy', 0)}%",
        f"Best Model: {ml.get('best_model', 'N/A')} ({ml.get('best_accuracy', 0)}%)",
        f"Data Health Score: {data['health']['score']}% ({data['health']['grade']})",
    ])

    # Slide 5: Recommendations
    s = add_slide()
    add_title(s, "Recommendations")
    add_bullets(s, [f"{c['title']}: {c['body']}" for c in data["recommendation_cards"]])

    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)

    out_name = f"{filename.rsplit('.', 1)[0]}_AI_Presentation.pptx"
    return Response(
        content=buf.read(),
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        headers={"Content-Disposition": f'attachment; filename="{out_name}"'},
    )
